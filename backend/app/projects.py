"""Claude 프로젝트 방식의 작업공간 — 이름·설명·지침·지식 파일을 가진 대화 묶음.

프로젝트에 속한 대화는 매 요청마다 (1) 프로젝트 지침과 (2) 업로드한 지식 파일 본문을
시스템 프롬프트로 받는다. 사업(사업현황) 연결은 선택 사항이라 걸어두면 그 사업의 현황
데이터가 추가로 붙는다(chat._프로젝트_컨텍스트)."""

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel

from backend.app import auth, repository as repo
from backend.app.files import (
    _hwp_텍스트_추출, _pdf_텍스트_추출, _업로드_원본_읽기, _파일버퍼,
)

router = APIRouter(dependencies=[Depends(auth.인증_확인)])

# 지식 용량 한도(글자 수). 한국어는 대략 글자당 0.6~1 토큰이라 150k자면 ~100k 토큰 안팎 —
# 매 요청에 전부 실리지만 시스템 프롬프트 캐시(ai_agent._시스템_프롬프트_구성)로 재사용된다.
지식_파일당_최대_글자수 = 60_000
지식_프로젝트당_최대_글자수 = 150_000
_지식_업로드_최대_바이트 = 10 * 1024 * 1024
지식_허용_확장자 = (".txt", ".md", ".csv", ".xlsx", ".xls", ".pdf", ".hwp", ".docx")
_지침_최대_글자수 = 8000


def 소유권_확인(프로젝트_id: int, 사용자_id: int) -> dict:
    """대화의 _소유권_확인과 같은 규칙 — 사용자_id가 NULL이면 레거시라 모두에게 열려있다."""
    프로젝트 = repo.프로젝트_조회(프로젝트_id)
    if not 프로젝트:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다.")
    if 프로젝트["사용자_id"] is not None and 프로젝트["사용자_id"] != 사용자_id:
        raise HTTPException(status_code=403, detail="다른 계정의 프로젝트입니다.")
    return 프로젝트


def 프로젝트_시스템_텍스트(프로젝트: dict | None) -> str:
    """프로젝트 지침 + 지식 파일 본문을 시스템 프롬프트에 덧붙일 한 덩어리로 만든다."""
    if not 프로젝트:
        return ""
    조각들 = [f"[현재 대화는 사용자의 프로젝트 '{프로젝트['이름']}'에 속해 있습니다.]"]
    if 프로젝트.get("설명"):
        조각들.append(f"프로젝트 설명: {프로젝트['설명']}")
    if (프로젝트.get("지침") or "").strip():
        조각들.append(
            "## 프로젝트 지침\n사용자가 이 프로젝트의 모든 대화에 적용하라고 지정한 지침입니다. "
            "위의 기본 지침과 충돌하지 않는 한 항상 따르세요.\n" + 프로젝트["지침"].strip()
        )
    지식들 = repo.프로젝트_지식_본문_불러오기(프로젝트["id"])
    if 지식들:
        본문 = "\n\n".join(f"### {k['파일명']}\n{k['내용']}" for k in 지식들)
        조각들.append(
            "## 프로젝트 지식\n사용자가 이 프로젝트에 올려둔 참고 자료입니다. 관련 질문이면 "
            "근거로 활용하고 어느 파일에서 왔는지 파일명을 언급하세요. 자료에 없는 내용은 "
            "있는 것처럼 지어내지 마세요. 사업현황 DB와 값이 다르면 둘 다 알려주세요.\n\n" + 본문
        )
    return "\n\n".join(조각들)


def 지식_텍스트_추출(파일명: str, 내용: bytes) -> str:
    """업로드된 지식 파일을 텍스트로 바꾼다. 실패/지원 불가 형식은 ValueError(사용자용 메시지)."""
    이름 = 파일명.lower()
    버퍼 = _파일버퍼(내용, 파일명)
    한도 = 지식_파일당_최대_글자수
    if 이름.endswith((".txt", ".md")):
        try:
            텍스트 = 내용.decode("utf-8-sig")
        except UnicodeDecodeError:
            텍스트 = 내용.decode("cp949", errors="ignore")
    elif 이름.endswith((".csv", ".xlsx", ".xls")):
        텍스트 = _업로드_원본_읽기(버퍼).to_csv(index=False)
    elif 이름.endswith(".pdf"):
        텍스트 = _pdf_텍스트_추출(버퍼, 최대글자수=한도)
    elif 이름.endswith(".hwp"):
        텍스트 = _hwp_텍스트_추출(버퍼, 최대글자수=한도)
    elif 이름.endswith(".docx"):
        from docx import Document

        문서 = Document(버퍼)
        줄들 = [p.text for p in 문서.paragraphs if p.text.strip()]
        for 표 in 문서.tables:
            for 행 in 표.rows:
                줄들.append(" | ".join(셀.text.strip() for 셀 in 행.cells))
        텍스트 = "\n".join(줄들)
    else:
        raise ValueError("지원하지 않는 형식이에요. (txt, md, csv, xlsx, pdf, hwp, docx)")

    텍스트 = 텍스트.strip()
    if not 텍스트:
        raise ValueError("파일에서 텍스트를 읽지 못했어요. (스캔 이미지로만 된 PDF는 지원하지 않아요)")
    if len(텍스트) > 한도:
        텍스트 = 텍스트[:한도] + "\n...(이하 생략)"
    return 텍스트


# ---------------- 엔드포인트 ----------------


class _프로젝트_생성_요청(BaseModel):
    이름: str
    설명: str = ""
    지침: str = ""
    사업_id: int | None = None


class _프로젝트_수정_요청(BaseModel):
    이름: str | None = None
    설명: str | None = None
    지침: str | None = None
    사업_id: int | None = None  # 명시적으로 null을 보내면 사업 연결 해제 — model_fields_set으로 구분


def _사업_라벨(사업_id: int | None) -> str | None:
    if not 사업_id:
        return None
    df = repo.사업현황_불러오기()
    행 = df[df["id"] == 사업_id] if "id" in df.columns else df.iloc[0:0]
    if 행.empty:
        return None
    r = 행.iloc[0]
    return f"{r.get('업체명', '') or ''} · {r.get('용역명', '') or ''}".strip(" ·")


def _이름_정리(값: str) -> str:
    이름 = (값 or "").strip()
    if not 이름:
        raise HTTPException(status_code=400, detail="프로젝트 이름을 입력해주세요.")
    return 이름[:60]


def _지침_검사(값: str) -> str:
    if len(값) > _지침_최대_글자수:
        raise HTTPException(status_code=400, detail=f"지침은 {_지침_최대_글자수:,}자까지 쓸 수 있어요.")
    return 값


@router.get("/api/projects")
def 프로젝트_목록(사용자: dict = Depends(auth.현재_사용자)):
    목록 = repo.프로젝트_목록_불러오기(사용자["id"])
    for p in 목록:
        p["사업_라벨"] = _사업_라벨(p["사업_id"])
    return 목록


@router.post("/api/projects")
def 프로젝트_생성(요청: _프로젝트_생성_요청, 사용자: dict = Depends(auth.현재_사용자)):
    프로젝트_id = repo.프로젝트_생성(
        사용자["id"], _이름_정리(요청.이름), (요청.설명 or "").strip()[:300],
        _지침_검사(요청.지침 or ""), 요청.사업_id,
    )
    return {"id": 프로젝트_id}


@router.get("/api/projects/{project_id}")
def 프로젝트_상세(project_id: int, 사용자: dict = Depends(auth.현재_사용자)):
    프로젝트 = 소유권_확인(project_id, 사용자["id"])
    대화들 = [c for c in repo.대화_목록_불러오기(사용자["id"]) if c.get("프로젝트_id") == project_id]
    return {
        "id": 프로젝트["id"],
        "이름": 프로젝트["이름"],
        "설명": 프로젝트["설명"] or "",
        "지침": 프로젝트["지침"] or "",
        "사업_id": 프로젝트["사업_id"],
        "사업_라벨": _사업_라벨(프로젝트["사업_id"]),
        "지식": repo.프로젝트_지식_목록(project_id),
        "지식_글자수": repo.프로젝트_지식_총글자수(project_id),
        "지식_한도": 지식_프로젝트당_최대_글자수,
        "대화": 대화들,
    }


@router.patch("/api/projects/{project_id}")
def 프로젝트_수정(project_id: int, 요청: _프로젝트_수정_요청, 사용자: dict = Depends(auth.현재_사용자)):
    소유권_확인(project_id, 사용자["id"])
    변경: dict = {}
    if "이름" in 요청.model_fields_set:
        변경["이름"] = _이름_정리(요청.이름 or "")
    if "설명" in 요청.model_fields_set:
        변경["설명"] = (요청.설명 or "").strip()[:300]
    if "지침" in 요청.model_fields_set:
        변경["지침"] = _지침_검사(요청.지침 or "")
    if "사업_id" in 요청.model_fields_set:
        변경["사업_id"] = 요청.사업_id
    repo.프로젝트_수정(project_id, 변경)
    return {"ok": True}


@router.delete("/api/projects/{project_id}")
def 프로젝트_삭제(project_id: int, 사용자: dict = Depends(auth.현재_사용자)):
    소유권_확인(project_id, 사용자["id"])
    repo.프로젝트_삭제(project_id)
    return {"ok": True}


@router.post("/api/projects/{project_id}/knowledge")
async def 지식_추가(project_id: int, file: UploadFile = File(...), 사용자: dict = Depends(auth.현재_사용자)):
    소유권_확인(project_id, 사용자["id"])
    파일명 = file.filename or "파일"
    if not 파일명.lower().endswith(지식_허용_확장자):
        raise HTTPException(status_code=400, detail="지원하지 않는 형식이에요. (txt, md, csv, xlsx, pdf, hwp, docx)")
    내용 = await file.read()
    if len(내용) > _지식_업로드_최대_바이트:
        raise HTTPException(status_code=400, detail="파일이 너무 커요. (최대 10MB)")
    try:
        텍스트 = 지식_텍스트_추출(파일명, 내용)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"파일을 읽지 못했어요: {e}")

    남은_글자수 = 지식_프로젝트당_최대_글자수 - repo.프로젝트_지식_총글자수(project_id)
    if len(텍스트) > 남은_글자수:
        raise HTTPException(
            status_code=400,
            detail=f"프로젝트 지식 용량을 넘어요. (이 파일 {len(텍스트):,}자 / 남은 용량 {max(남은_글자수, 0):,}자) "
                   "다른 파일을 지우고 다시 올려주세요.",
        )
    지식_id = repo.프로젝트_지식_추가(project_id, 파일명, 텍스트)
    return {"id": 지식_id, "글자수": len(텍스트)}


@router.delete("/api/projects/{project_id}/knowledge/{knowledge_id}")
def 지식_삭제(project_id: int, knowledge_id: int, 사용자: dict = Depends(auth.현재_사용자)):
    소유권_확인(project_id, 사용자["id"])
    if not repo.프로젝트_지식_삭제(project_id, knowledge_id):
        raise HTTPException(status_code=404, detail="지식 파일을 찾을 수 없습니다.")
    return {"ok": True}
