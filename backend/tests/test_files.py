import io
import zipfile
from pathlib import Path

import pytest

from backend.app import files

# pyhwp(https://github.com/mete0r/pyhwp, AGPLv3+) 자체 테스트 스위트의 픽스처를 그대로
# 가져왔다 — 실제 바이너리 .hwp를 코드로 만들 방법이 없어, 진짜 한글 문서로 회귀를 검증하려면
# 이 방법뿐이다. charshape.hwp는 문단 텍스트(한글·영문·아랍어 혼용) 검증용,
# borderfill.hwp는 표 구조(칸이 비어 있음) 검증용.
_픽스처_경로 = Path(__file__).parent / "fixtures"


def test_docx_텍스트_추출():
    from docx import Document

    문서 = Document()
    문서.add_paragraph("첫 번째 문단입니다.")
    표 = 문서.add_table(rows=1, cols=2)
    표.rows[0].cells[0].text = "이름"
    표.rows[0].cells[1].text = "값"
    버퍼 = io.BytesIO()
    문서.save(버퍼)

    텍스트 = files._docx_텍스트_추출(files._파일버퍼(버퍼.getvalue(), "테스트.docx"))
    assert "첫 번째 문단입니다." in 텍스트
    assert "이름 | 값" in 텍스트


def test_pptx_텍스트_추출():
    from pptx import Presentation
    from pptx.util import Inches

    프레젠테이션 = Presentation()
    빈_레이아웃 = 프레젠테이션.slide_layouts[6]
    슬라이드 = 프레젠테이션.slides.add_slide(빈_레이아웃)
    텍스트상자 = 슬라이드.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    텍스트상자.text_frame.text = "슬라이드 본문입니다."
    버퍼 = io.BytesIO()
    프레젠테이션.save(버퍼)

    텍스트 = files._pptx_텍스트_추출(files._파일버퍼(버퍼.getvalue(), "테스트.pptx"))
    assert "슬라이드 본문입니다." in 텍스트
    assert "[슬라이드 1]" in 텍스트


def _최소_hwpx_바이트(본문_xml: str) -> bytes:
    버퍼 = io.BytesIO()
    with zipfile.ZipFile(버퍼, "w") as zf:
        zf.writestr("Contents/section0.xml", 본문_xml)
    return 버퍼.getvalue()


def test_hwpx_텍스트_추출():
    section_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<hs:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
        'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">'
        "<hp:p><hp:run><hp:t>테스트 문단입니다.</hp:t></hp:run></hp:p>"
        "</hs:sec>"
    )
    텍스트 = files._hwpx_텍스트_추출(files._파일버퍼(_최소_hwpx_바이트(section_xml), "테스트.hwpx"))
    assert 텍스트 == "테스트 문단입니다."


def test_hwpx_텍스트_추출_여러_문단_순서_보존():
    section_xml = (
        '<hs:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
        'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">'
        "<hp:p><hp:run><hp:t>첫 문단</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:t>둘째 </hp:t></hp:run><hp:run><hp:t>문단</hp:t></hp:run></hp:p>"
        "</hs:sec>"
    )
    텍스트 = files._hwpx_텍스트_추출(files._파일버퍼(_최소_hwpx_바이트(section_xml), "테스트.hwpx"))
    assert 텍스트 == "첫 문단\n둘째 문단"


def test_hwpx_텍스트_추출_섹션_없으면_오류():
    버퍼 = io.BytesIO()
    with zipfile.ZipFile(버퍼, "w") as zf:
        zf.writestr("META-INF/manifest.xml", "<x/>")
    with pytest.raises(ValueError):
        files._hwpx_텍스트_추출(files._파일버퍼(버퍼.getvalue(), "빈.hwpx"))


def test_hwp_텍스트_추출_실제_문서():
    """예전(hwp5txt) 방식은 표·그림을 <표>/<그림> placeholder로만 남겼다 — hwp5html 기반으로
    바꾼 뒤 실제 문단 텍스트가 순서대로 나오는지 진짜 .hwp 픽스처로 확인한다."""
    바이트 = (_픽스처_경로 / "charshape.hwp").read_bytes()
    텍스트 = files._hwp_텍스트_추출(files._파일버퍼(바이트, "charshape.hwp"))
    assert "글자크기" in 텍스트
    assert "<표>" not in 텍스트 and "<그림>" not in 텍스트


def test_hwp_텍스트_추출_표_구조_있는_문서는_오류_없이_처리():
    """표 칸이 비어 있는 픽스처 — 표 자체가 있어도 예외 없이(빈 결과로) 처리되는지만 확인."""
    바이트 = (_픽스처_경로 / "borderfill.hwp").read_bytes()
    텍스트 = files._hwp_텍스트_추출(files._파일버퍼(바이트, "borderfill.hwp"))
    assert isinstance(텍스트, str)
